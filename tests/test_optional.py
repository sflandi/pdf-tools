"""Features that need programs outside Python. Skipped when the program is not installed."""

import os

import pymupdf
import pytest

from pdftools import ocr, office


@pytest.mark.skipif(not ocr.available_languages(), reason="Tesseract language data not installed")
def test_ocr_makes_scan_searchable(tmp_path):
    src = pymupdf.open()
    page = src.new_page()
    page.insert_text((72, 120), "Annual Financial Report", fontsize=28)
    scan = pymupdf.open()
    pix = page.get_pixmap(dpi=200)
    sp = scan.new_page(width=page.rect.width, height=page.rect.height)
    sp.insert_image(sp.rect, pixmap=pix)  # the page is only an image, like a scan
    path = str(tmp_path / "scan.pdf")
    scan.save(path)
    out = str(tmp_path / "ocr.pdf")
    lang = ocr.available_languages()[0]
    assert ocr.ocr_to_searchable_pdf(path, out, lang) == 1
    assert "Financial" in pymupdf.open(out)[0].get_text()


@pytest.mark.skipif(not office.find_soffice(), reason="LibreOffice not installed")
def test_office_to_pdf(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("Halo dari Word")
    src = str(tmp_path / "a.docx")
    d.save(src)
    pdf = office.office_to_pdf(src, str(tmp_path))
    assert "Halo dari Word" in pymupdf.open(pdf)[0].get_text()


def test_pdf2docx_worker(text_pdf, tmp_path):
    from pdftools import ui

    out = str(tmp_path / "x.docx")
    ui.run_worker(["pdf2docx", text_pdf, out, "--pages", "0,2"], timeout=120)
    from docx import Document

    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Page 1" in text and "Page 3" in text and "Page 2" not in text
